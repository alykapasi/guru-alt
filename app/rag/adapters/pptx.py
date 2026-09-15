"""PPTX adapter — one located unit per slide (python-pptx).

**Tables and grouped shapes are read, and until this they were not (S27).** The walk asked each
shape for a text frame, and a table is a ``GraphicFrame`` that has none — so every table on
every slide was absent from ingestion, exactly as Word's were. A group is the same story from
the other direction: it holds the shapes, and the shapes hold the text, so anything a deck
author had grouped disappeared with it. Neither failed loudly; the slide simply arrived with
less on it than it had.
"""

from pathlib import Path

from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape

from app.rag.adapters.base import ExtractContext, ExtractedUnit, table_text

_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _shape_text(shape: BaseShape) -> str:
    """Whatever this shape carries, including through a group.

    Order follows the shape order the deck stores, so a table stays where its heading put it.
    """
    # A group holds no text of its own; everything is in what it contains.
    if isinstance(shape, GroupShape):
        return "\n".join(
            text for text in (_shape_text(child) for child in shape.shapes) if text.strip()
        )
    if isinstance(shape, GraphicFrame) and shape.has_table:
        return table_text([cell.text for cell in row.cells] for row in shape.table.rows)
    # `getattr` rather than a third isinstance: several unrelated classes carry a text frame
    # (autoshapes, placeholders, text boxes) and naming them all would be a list to keep in
    # step with a library we do not control.
    frame = getattr(shape, "text_frame", None)
    return str(frame.text) if frame is not None else ""


class PptxAdapter:
    name = "pptx"

    def handles(self, content_type: str) -> bool:
        return content_type == _CONTENT_TYPE

    async def extract(self, path: Path, *, meta: dict, ctx: ExtractContext) -> list[ExtractedUnit]:
        presentation = Presentation(str(path))
        units: list[ExtractedUnit] = []
        for index, slide in enumerate(presentation.slides):
            lines = [text for text in (_shape_text(s) for s in slide.shapes) if text.strip()]
            text = "\n".join(lines)
            if text.strip():
                units.append(ExtractedUnit(text=text, locator={"slide": index + 1}))
        return units
