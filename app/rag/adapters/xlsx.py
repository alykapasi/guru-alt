"""XLSX adapter — one located unit per worksheet (openpyxl).

A spreadsheet is a table and almost nothing else, so how a row is rendered is the whole of what
gets indexed. Two things used to destroy it before chunking ever saw it (S27): cells were joined
with a space, so nothing marked where one column ended and the next began, and empty cells were
dropped entirely, so a row with a gap in the middle came out with its later values shifted left
into the wrong columns — indistinguishable from a shorter row. A learner asking about the Q2
figure could be answered with Q3's.

Rows are tab-separated now, which is the separator normalization keeps as a column boundary
(see ``app.rag.chunking._collapse``), and an empty cell is an empty field rather than a missing
one.
"""

from pathlib import Path

from openpyxl import load_workbook

from app.rag.adapters.base import ExtractContext, ExtractedUnit

_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_COLUMN = "\t"


class XlsxAdapter:
    name = "xlsx"

    def handles(self, content_type: str) -> bool:
        return content_type == _CONTENT_TYPE

    async def extract(self, path: Path, *, meta: dict, ctx: ExtractContext) -> list[ExtractedUnit]:
        # Pass a file handle (not the path) so openpyxl skips its extension check — our temp
        # file has no .xlsx suffix. read_only is lazy, so keep the handle open while iterating.
        with open(path, "rb") as handle:
            workbook = load_workbook(handle, read_only=True, data_only=True)
            try:
                units: list[ExtractedUnit] = []
                for worksheet in workbook.worksheets:
                    rows = [
                        _COLUMN.join("" if cell is None else str(cell) for cell in row)
                        for row in worksheet.iter_rows(values_only=True)
                    ]
                    # A row of nothing but separators is an empty row, and dropping it is not
                    # the same as dropping an empty cell: no column is displaced by it.
                    text = "\n".join(row for row in rows if row.strip(_COLUMN).strip())
                    if text.strip():
                        units.append(ExtractedUnit(text=text, locator={"sheet": worksheet.title}))
                return units
            finally:
                workbook.close()
