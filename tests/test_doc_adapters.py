"""Per-format doc adapters: extraction + locators, and PDF dispatch through the pipeline."""

import asyncio
import io
import os
import tempfile
import uuid
from collections.abc import Sequence
from pathlib import Path

import pymupdf
from docx import Document
from ebooklib import epub
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.llm import ChatMessage, ChatResponse, ModelRole, ToolDef
from app.llm.providers import FakeProvider
from app.llm.registry import LLMClient, ModelSpec, fake_llm_client
from app.models.learner import Learner
from app.models.source import Chunk, SourceKind, SourceStatus
from app.rag.adapters import ExtractContext, select_adapter
from app.rag.adapters.docx import DocxAdapter
from app.rag.adapters.epub import EpubAdapter
from app.rag.adapters.pdf import PdfAdapter
from app.rag.adapters.pptx import PptxAdapter
from app.rag.adapters.xlsx import XlsxAdapter
from app.services import ingestion
from app.storage import InMemoryBlobStore

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --- fixture builders (generate real files in-memory) -----------------------


def _png() -> bytes:
    """A small solid-colour PNG that PyMuPDF can embed (has no text)."""
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 16, 16))
    pix.clear_with(220)
    return pix.tobytes("png")


_PNG = _png()


def _path(tmp_path: Path, data: bytes, name: str = "f.bin") -> Path:
    """Write bytes to a temp file and return its path (adapters now read from a path)."""
    p = tmp_path / name
    p.write_bytes(data)
    return p


def _pdf(pages: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _image_pdf(n_pages: int = 1) -> bytes:
    """A PDF whose pages carry an image and no extractable text (a 'scanned' doc)."""
    doc = pymupdf.open()
    for _ in range(n_pages):
        page = doc.new_page()
        page.insert_image(page.rect, stream=_PNG)
    data = doc.tobytes()
    doc.close()
    return data


def _mixed_pdf() -> bytes:
    """Page 1 is born-digital text; page 2 is an image-only 'scanned' page."""
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Born-digital text on the first page of the document")
    img_page = doc.new_page()
    img_page.insert_image(img_page.rect, stream=_PNG)
    data = doc.tobytes()
    doc.close()
    return data


def _docx(paragraphs: list[str]) -> bytes:
    document = Document()
    for para in paragraphs:
        document.add_paragraph(para)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _docx_with_table(intro: str, rows: list[list[str]], outro: str) -> bytes:
    """A document whose table sits between two paragraphs, which is where reports put them."""
    document = Document()
    document.add_paragraph(intro)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph(outro)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pptx(slides: list[str]) -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for text in slides:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xlsx(title: str, rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- extraction -------------------------------------------------------------


async def test_pdf_adapter_extracts_one_unit_per_page(tmp_path: Path) -> None:
    units = await PdfAdapter().extract(
        _path(tmp_path, _pdf(["Hello from page one", "Second page content"])),
        meta={},
        ctx=ExtractContext(),
    )
    assert len(units) == 2
    assert units[0].locator == {"page": 1} and "page one" in units[0].text
    assert units[1].locator == {"page": 2} and "Second page" in units[1].text


async def test_docx_adapter_joins_paragraphs(tmp_path: Path) -> None:
    units = await DocxAdapter().extract(
        _path(tmp_path, _docx(["First paragraph.", "Second paragraph."])),
        meta={},
        ctx=ExtractContext(),
    )
    assert len(units) == 1
    assert "First paragraph" in units[0].text and "Second paragraph" in units[0].text


async def test_docx_adapter_reads_tables_it_used_to_drop_entirely(tmp_path: Path) -> None:
    """`document.paragraphs` returns only the body's top-level paragraphs, so every cell of
    every table was absent from ingestion — not mangled, gone, with the source reported `done`
    and nothing anywhere saying the figures never arrived."""
    units = await DocxAdapter().extract(
        _path(
            tmp_path,
            _docx_with_table(
                "Quarterly results follow.",
                [["Region", "Q1", "Q2"], ["North", "1200", "1450"]],
                "End of report.",
            ),
        ),
        meta={},
        ctx=ExtractContext(),
    )

    assert len(units) == 1
    assert "1450" in units[0].text


async def test_docx_table_rows_keep_their_columns_and_their_place(tmp_path: Path) -> None:
    """Order matters as much as presence: a table belongs under the paragraph that introduced
    it, and appending every table at the end would file them all under the last one."""
    units = await DocxAdapter().extract(
        _path(
            tmp_path,
            _docx_with_table(
                "Quarterly results follow.",
                [["Region", "Q1", "Q2"], ["North", "1200", "1450"]],
                "End of report.",
            ),
        ),
        meta={},
        ctx=ExtractContext(),
    )

    assert units[0].text.splitlines() == [
        "Quarterly results follow.",
        "Region\tQ1\tQ2",
        "North\t1200\t1450",
        "End of report.",
    ]


async def test_docx_keeps_an_empty_cell_rather_than_shifting_the_row(tmp_path: Path) -> None:
    """A dropped empty cell moves every later value one column left, which turns a gap into
    wrong data — the row still parses, and Q2's figure now sits under Q1."""
    units = await DocxAdapter().extract(
        _path(
            tmp_path,
            _docx_with_table("Results", [["Region", "Q1", "Q2"], ["North", "", "1450"]], "End."),
        ),
        meta={},
        ctx=ExtractContext(),
    )

    assert "North\t\t1450" in units[0].text


async def test_docx_keeps_a_multi_paragraph_cell_on_one_row(tmp_path: Path) -> None:
    """A cell can hold several paragraphs, and `cell.text` joins them with a newline. Left
    alone that newline ends the row early: the rest of the cell becomes its own line with no
    columns at all, and every column after it in the real row is orphaned."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Note"
    table.cell(1, 0).text = "North"
    cell = table.cell(1, 1)
    cell.text = "first line"
    cell.add_paragraph("second line")
    buf = io.BytesIO()
    document.save(buf)

    units = await DocxAdapter().extract(
        _path(tmp_path, buf.getvalue()), meta={}, ctx=ExtractContext()
    )

    assert units[0].text.splitlines() == [
        "Region\tNote",
        "North\tfirst line second line",
    ]


async def test_pptx_adapter_extracts_one_unit_per_slide(tmp_path: Path) -> None:
    units = await PptxAdapter().extract(
        _path(tmp_path, _pptx(["Slide one bullet", "Slide two bullet"])),
        meta={},
        ctx=ExtractContext(),
    )
    assert len(units) == 2
    assert units[0].locator == {"slide": 1} and "Slide one" in units[0].text
    assert units[1].locator == {"slide": 2}


def _pptx_with_table(heading: str, rows: list[list[str]]) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(1))
    box.text_frame.text = heading
    frame = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(1), Inches(2), Inches(4), Inches(1)
    )
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            frame.table.cell(r, c).text = value
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def test_pptx_reads_a_table_it_used_to_skip_entirely(tmp_path: Path) -> None:
    """A table is a `GraphicFrame`, which has no text frame — so the shape walk skipped it and
    every table on every slide was absent from ingestion, the same hole Word had."""
    units = await PptxAdapter().extract(
        _path(
            tmp_path, _pptx_with_table("Quarterly results", [["Region", "Q1"], ["North", "1200"]])
        ),
        meta={},
        ctx=ExtractContext(),
    )

    assert units[0].text.splitlines() == ["Quarterly results", "Region\tQ1", "North\t1200"]


async def test_pptx_reads_through_a_group(tmp_path: Path) -> None:
    """A group holds the shapes and the shapes hold the text, so anything a deck author had
    grouped disappeared with it — silently, as a slide with less on it than it had."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    first.text_frame.text = "Inside a group"
    second = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    second.text_frame.text = "Also grouped"
    slide.shapes.add_group_shape([first, second])
    buf = io.BytesIO()
    prs.save(buf)

    units = await PptxAdapter().extract(
        _path(tmp_path, buf.getvalue()), meta={}, ctx=ExtractContext()
    )

    assert "Inside a group" in units[0].text
    assert "Also grouped" in units[0].text


async def test_xlsx_adapter_extracts_one_unit_per_sheet(tmp_path: Path) -> None:
    units = await XlsxAdapter().extract(
        _path(tmp_path, _xlsx("Scores", [["Name", "Score"], ["Alice", 90]])),
        meta={},
        ctx=ExtractContext(),
    )
    assert len(units) == 1
    assert units[0].locator == {"sheet": "Scores"}
    assert "Alice" in units[0].text and "90" in units[0].text


async def test_xlsx_rows_are_tab_separated_so_the_columns_survive(tmp_path: Path) -> None:
    """A spreadsheet is a table and almost nothing else. Joining the cells with a space left
    nothing to say where one column ended, which is how a table became a run of numbers."""
    units = await XlsxAdapter().extract(
        _path(tmp_path, _xlsx("Q", [["Region", "Q1", "Q2"], ["North", 1200, 1450]])),
        meta={},
        ctx=ExtractContext(),
    )

    assert units[0].text.splitlines() == ["Region\tQ1\tQ2", "North\t1200\t1450"]


async def test_xlsx_keeps_an_empty_cell_rather_than_shifting_the_row(tmp_path: Path) -> None:
    """`if cell is not None` dropped the gap and moved every later value left, so a learner
    asking about Q2 could be answered with Q3's figure."""
    units = await XlsxAdapter().extract(
        _path(tmp_path, _xlsx("Q", [["Region", "Q1", "Q2"], ["North", None, 1450]])),
        meta={},
        ctx=ExtractContext(),
    )

    assert units[0].text.splitlines()[1] == "North\t\t1450"


async def test_xlsx_drops_a_blank_row_because_it_displaces_no_column(tmp_path: Path) -> None:
    units = await XlsxAdapter().extract(
        _path(tmp_path, _xlsx("Q", [["Region", "Q1"], [None, None], ["North", 1200]])),
        meta={},
        ctx=ExtractContext(),
    )

    assert units[0].text.splitlines() == ["Region\tQ1", "North\t1200"]


def test_registry_dispatches_by_content_type() -> None:
    assert isinstance(select_adapter("application/pdf"), PdfAdapter)
    assert isinstance(select_adapter(DOCX_CT), DocxAdapter)
    assert isinstance(select_adapter(PPTX_CT), PptxAdapter)
    assert isinstance(select_adapter(XLSX_CT), XlsxAdapter)
    assert select_adapter("application/x-unknown") is None


# --- pipeline dispatch ------------------------------------------------------


async def test_ingest_pdf_records_page_provenance(db_session: AsyncSession) -> None:
    store = InMemoryBlobStore()
    learner = Learner(handle=f"l-{uuid.uuid4().hex[:8]}")
    db_session.add(learner)
    await db_session.flush()

    data = _pdf(["Photosynthesis occurs in chloroplasts.", "Respiration occurs in mitochondria."])
    source = await ingestion.create_source(
        db_session,
        store,
        learner_id=learner.id,
        kind=SourceKind.FILE,
        origin="biology.pdf",
        content_type="application/pdf",
        data=data,
    )
    result = await ingestion.ingest_source(db_session, store, fake_llm_client(), source.id)
    assert result is not None  # the source was claimable
    assert result.status == SourceStatus.DONE

    chunks = (
        await db_session.scalars(
            select(Chunk).where(Chunk.source_id == source.id).order_by(Chunk.ordinal)
        )
    ).all()
    assert len(chunks) >= 2
    assert chunks[0].provenance["method"] == "pdf"
    assert {c.provenance["page"] for c in chunks} == {1, 2}


async def test_an_ingested_spreadsheet_is_stored_with_its_columns(
    db_session: AsyncSession,
) -> None:
    """The claim end to end, on the text that is actually stored. The adapter and the chunker
    each had to stop destroying the table, and asserting on either alone would leave the other
    free to undo it — which is how it stood before: the adapter joined cells with a space, and
    normalization would have flattened a tab anyway."""
    store = InMemoryBlobStore()
    learner = Learner(handle=f"l-{uuid.uuid4().hex[:8]}")
    db_session.add(learner)
    await db_session.flush()

    source = await ingestion.create_source(
        db_session,
        store,
        learner_id=learner.id,
        kind=SourceKind.FILE,
        origin="results.xlsx",
        content_type=XLSX_CT,
        data=_xlsx("Q", [["Region", "Q1", "Q2"], ["North", None, 1450], ["South", 980, 1010]]),
    )
    await ingestion.ingest_source(db_session, store, fake_llm_client(), source.id)

    chunk = await db_session.scalar(select(Chunk).where(Chunk.source_id == source.id))
    assert chunk is not None
    assert chunk.text.splitlines() == [
        "Region\tQ1\tQ2",
        "North\t\t1450",
        "South\t980\t1010",
    ]


# --- scanned-PDF OCR fallback -----------------------------------------------


async def test_pdf_born_digital_skips_ocr(tmp_path: Path) -> None:
    ctx = ExtractContext(llm=fake_llm_client("should not be called"))
    units = await PdfAdapter().extract(
        _path(tmp_path, _pdf(["Plenty of real digital text on this page"])), meta={}, ctx=ctx
    )
    assert len(units) == 1
    assert units[0].method is None  # text path, not OCR
    assert ctx.usage_log == []  # the vision model was never called


async def test_pdf_scanned_page_is_ocred(tmp_path: Path) -> None:
    ctx = ExtractContext(llm=fake_llm_client("handwritten lecture notes"))
    units = await PdfAdapter().extract(_path(tmp_path, _image_pdf(1)), meta={}, ctx=ctx)
    assert len(units) == 1
    assert units[0].method == "ocr"
    assert units[0].text == "handwritten lecture notes"
    assert units[0].locator == {"page": 1}
    assert len(ctx.usage_log) == 1  # one vision call for the scanned page


async def test_pdf_mixed_pages_use_per_page_method(tmp_path: Path) -> None:
    ctx = ExtractContext(llm=fake_llm_client("scanned text"))
    units = await PdfAdapter().extract(_path(tmp_path, _mixed_pdf()), meta={}, ctx=ctx)
    assert len(units) == 2
    assert units[0].method is None and units[0].locator == {"page": 1}  # text page
    assert units[1].method == "ocr" and units[1].locator == {"page": 2}  # scanned page
    assert len(ctx.usage_log) == 1  # only the scanned page hit vision


async def test_pdf_scanned_page_without_llm_is_skipped(tmp_path: Path) -> None:
    units = await PdfAdapter().extract(
        _path(tmp_path, _image_pdf(1)), meta={}, ctx=ExtractContext()
    )
    assert units == []  # no text and no OCR capability -> nothing extracted


async def test_ingest_scanned_pdf_records_ocr_provenance(db_session: AsyncSession) -> None:
    store = InMemoryBlobStore()
    learner = Learner(handle=f"l-{uuid.uuid4().hex[:8]}")
    db_session.add(learner)
    await db_session.flush()

    source = await ingestion.create_source(
        db_session,
        store,
        learner_id=learner.id,
        kind=SourceKind.FILE,
        origin="scan.pdf",
        content_type="application/pdf",
        data=_image_pdf(2),
    )
    result = await ingestion.ingest_source(
        db_session, store, fake_llm_client("scanned page text"), source.id
    )
    assert result is not None  # the source was claimable
    assert result.status == SourceStatus.DONE

    chunks = (
        await db_session.scalars(
            select(Chunk).where(Chunk.source_id == source.id).order_by(Chunk.ordinal)
        )
    ).all()
    assert len(chunks) == 2
    assert all(c.provenance["method"] == "ocr" for c in chunks)
    assert {c.provenance["page"] for c in chunks} == {1, 2}


# --- concurrent scanned-PDF OCR (Slice A3) ----------------------------------


class _CountingVisionProvider(FakeProvider):
    """A fake vision provider that records the peak number of concurrent OCR calls."""

    def __init__(self, reply: str = "ocr text") -> None:
        super().__init__(reply=reply)
        self.inflight = 0
        self.max_inflight = 0

    async def complete(
        self,
        *,
        model: str,
        messages: Sequence[ChatMessage],
        system: str | None = None,
        max_tokens: int = 1024,
        tools: Sequence[ToolDef] | None = None,
    ) -> ChatResponse:
        self.inflight += 1
        self.max_inflight = max(self.max_inflight, self.inflight)
        try:
            await asyncio.sleep(0.02)  # hold the slot so genuine overlap is observable
            return await super().complete(
                model=model, messages=messages, system=system, max_tokens=max_tokens
            )
        finally:
            self.inflight -= 1


def _counting_vision_client() -> tuple[_CountingVisionProvider, LLMClient]:
    provider = _CountingVisionProvider()
    client = LLMClient({"fake": provider}, {r: ModelSpec("fake", "fake-1") for r in ModelRole})
    return provider, client


async def test_pdf_scanned_pages_ocr_concurrently_bounded(tmp_path: Path) -> None:
    provider, client = _counting_vision_client()
    ctx = ExtractContext(llm=client, ocr_concurrency=3)

    units = await PdfAdapter().extract(_path(tmp_path, _image_pdf(6)), meta={}, ctx=ctx)

    # Every page transcribed, in page order (order preserved despite out-of-order completion).
    assert [u.locator["page"] for u in units] == [1, 2, 3, 4, 5, 6]
    assert all(u.method == "ocr" for u in units)
    assert len(ctx.usage_log) == 6
    # Concurrent, but never more than the limit in flight at once.
    assert 1 < provider.max_inflight <= 3


# --- EPUB adapter (Slice A4) -------------------------------------------------


def _epub(chapters: list[tuple[str, str]]) -> bytes:
    """A minimal valid EPUB — one spine chapter per (title, body), with a nav doc first."""
    book = epub.EpubBook()
    book.set_identifier("id-test")
    book.set_title("Test Book")
    book.set_language("en")
    items = []
    for i, (title, body) in enumerate(chapters):
        chapter = epub.EpubHtml(title=title, file_name=f"chap_{i}.xhtml", lang="en")
        chapter.content = f"<html><body><h1>{title}</h1><p>{body}</p></body></html>"
        book.add_item(chapter)
        items.append(chapter)
    book.toc = list(items)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *items]  # realistic: the nav document leads the spine
    fd, name = tempfile.mkstemp(suffix=".epub")
    os.close(fd)
    try:
        epub.write_epub(name, book)
        return Path(name).read_bytes()
    finally:
        os.unlink(name)


async def test_epub_adapter_extracts_one_unit_per_chapter(tmp_path: Path) -> None:
    data = _epub([("Chapter One", "The first chapter body."), ("Chapter Two", "Second body.")])
    units = await EpubAdapter().extract(
        _path(tmp_path, data, "book.epub"), meta={}, ctx=ExtractContext()
    )

    assert len(units) == 2  # the nav document is skipped, not counted as a chapter
    assert [u.locator["chapter"] for u in units] == [1, 2]
    assert [u.locator["title"] for u in units] == ["Chapter One", "Chapter Two"]
    assert "first chapter body" in units[0].text
    assert units[0].method is None  # falls back to the adapter's own name ("epub")


def test_registry_dispatches_epub() -> None:
    assert isinstance(select_adapter("application/epub+zip"), EpubAdapter)


async def test_ingest_epub_records_chapter_provenance(db_session: AsyncSession) -> None:
    store = InMemoryBlobStore()
    learner = Learner(handle=f"l-{uuid.uuid4().hex[:8]}")
    db_session.add(learner)
    await db_session.flush()

    source = await ingestion.create_source(
        db_session,
        store,
        learner_id=learner.id,
        kind=SourceKind.FILE,
        origin="book.epub",
        content_type="application/epub+zip",
        data=_epub([("Intro", "Welcome to the book."), ("Body", "The main material here.")]),
    )
    result = await ingestion.ingest_source(db_session, store, fake_llm_client(), source.id)
    assert result is not None  # the source was claimable
    assert result.status == SourceStatus.DONE

    chunks = (
        await db_session.scalars(
            select(Chunk).where(Chunk.source_id == source.id).order_by(Chunk.ordinal)
        )
    ).all()
    assert chunks
    assert all(c.provenance["method"] == "epub" for c in chunks)
    assert {"Intro", "Body"} <= {c.provenance["title"] for c in chunks}
